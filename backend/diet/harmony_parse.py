# -*- coding: utf-8 -*-
"""요양원 주간 식단표(Processed_*.xlsx) 폴더 → 조화 모델 학습용 JSON

  엑셀 형식 (정성노인·포천·은화케어 공통): 시트 menu · nutrient · category · ingredient
    menu       1행 = 구분 + 날짜 7개, 이후 '조식/중식/석식' 6행씩 (밥 · 국 · 주찬 · 부찬1 · 부찬2 · 김치 순서)
               간식 행(오전간식·오후간식)은 쓰지 않음
    nutrient   food_id · Menu · 영양소
    category   Menu · Category
    ingredient food_id · Menu · ingredient_id · Amount_g · Ingredient

  실행:  cd backend
         venv-win\\Scripts\\python -m diet.harmony_parse "<엑셀 폴더>" 포천 diet/data/harmony/pocheon_menus.json
"""
from __future__ import annotations

import datetime as dt
import glob
import json
import os
import sys

SLOTS = ("밥", "국", "주찬", "부찬1", "부찬2", "김치")
MEALS = {"조식": "아침", "중식": "점심", "석식": "저녁"}
NUTR = {"에너지(kcal)": "energy", "단백질(g)": "protein", "지방(g)": "fat", "탄수화물(g)": "carb",
        "나트륨(mg)": "sodium", "칼슘(mg)": "ca", "비타민 D(μg)": "vit_d", "식이섬유(g)": "fiber"}


def _date(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 30000 < v < 80000:          # 엑셀 날짜 일련번호 (예: 45319)
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v or "").strip()
    try:
        return dt.date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        return None


def _s(v):
    s = str(v).strip() if v is not None else ""
    return s or None


def parse_file(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["menu"].iter_rows(values_only=True))
    dates = [_date(v) for v in rows[0][1:]]
    meals = []
    for lab, meal in MEALS.items():
        block = [r for r in rows[1:] if _s(r[0]) == lab]
        for j, d in enumerate(dates, start=1):
            if not d:
                continue
            m = {"date": d, "meal": meal}
            for slot, r in zip(SLOTS, block):
                m[slot] = _s(r[j]) if j < len(r) else None
            meals.append(m)
    menus, cats, ings = {}, {}, {}
    if "nutrient" in wb.sheetnames:
        it = wb["nutrient"].iter_rows(values_only=True)
        head = [_s(h) for h in next(it)]
        for r in it:
            row = dict(zip(head, r))
            name = _s(row.get("Menu"))
            if name:
                menus[name] = {"food_id": _s(row.get("food_id")),
                               **{k: row.get(h) for h, k in NUTR.items() if h in row}}
    if "category" in wb.sheetnames:
        for r in list(wb["category"].iter_rows(values_only=True))[1:]:
            if _s(r[0]):
                cats[_s(r[0])] = _s(r[1])
    if "ingredient" in wb.sheetnames:
        it = wb["ingredient"].iter_rows(values_only=True)
        head = [_s(h) for h in next(it)]
        for r in it:
            row = dict(zip(head, r))
            name, ing = _s(row.get("Menu")), _s(row.get("Ingredient"))
            if name and ing:
                g = row.get("Amount_g")
                try:
                    g = float(g)
                except (TypeError, ValueError):
                    g = None
                ings.setdefault(name, {})[ing] = {"g": g, "id": _s(row.get("ingredient_id"))}
    return meals, menus, cats, ings


def parse_folder(folder, facility):
    by, menus, cats, ings = {}, {}, {}, {}
    files = sorted(glob.glob(os.path.join(folder, "*.xlsx")))
    for f in files:
        if os.path.basename(f).startswith("~$"):
            continue
        ms, mn, ct, ig = parse_file(f)
        for m in ms:
            by[(m["date"], m["meal"])] = m          # 겹치는 날짜는 뒤 파일 기준
        menus.update(mn); cats.update(ct)
        for name, d_ in ig.items():                    # 같은 메뉴의 재료는 파일끼리 합침
            ings.setdefault(name, {}).update(d_)
    order = {"아침": 0, "점심": 1, "저녁": 2}
    meals = sorted(by.values(), key=lambda m: (m["date"], order[m["meal"]]))
    return {"facility": facility, "files": len(files), "meals": meals, "menus": menus,
            "categories": cats, "ingredients": ings}


def main(folder, facility, out):
    d = parse_folder(folder, facility)
    ok = [m for m in d["meals"] if all(m.get(s) for s in ("국", "주찬", "부찬1", "부찬2"))]
    json.dump(d, open(out, "w", encoding="utf-8"), ensure_ascii=False)
    ds = [m["date"] for m in d["meals"]]
    print(f"{facility}: 파일 {d['files']}개 · 끼니 {len(d['meals'])} (국·주찬·부찬 4칸 모두 {len(ok)}) · "
          f"메뉴 {len(d['menus'])} · 기간 {min(ds)} ~ {max(ds)} → {out}")
    return d


if __name__ == "__main__":
    main(*sys.argv[1:4])
