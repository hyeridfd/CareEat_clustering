# -*- coding: utf-8 -*-
"""식사형태 9종 × 5일 × 아침·점심·저녁 제공 영양량 (잔반 반영 전) → 엑셀"""
import sys
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import nutrition_core as nc

KEY = ['총량(g)', '에너지(kcal)', '탄수화물(g)', '단백질(g)', '지질(g)', '총 식이섬유(g)', '총당류(g)',
       '칼슘(mg)', '철(mg)', '인(mg)', '칼륨(mg)', '나트륨(mg)', '아연(mg)', '마그네슘(mg)',
       '비타민 A (RAE)(μg RAE)', '비타민 D(μg)', '비타민 E(mg)', '비타민 C(mg)', '티아민(mg)',
       '리보플라빈(mg)', '나이아신(NE)(mg)', '비타민 B6(mg)', '엽산(μg DFE)', '비타민 B12(μg)',
       '콜레스테롤(mg)', '총 포화지방산 (g)', '오메가3지방산(g)', '오메가6지방산(g)', '수분(g)']

ARIAL, HDR = 'Arial', PatternFill('solid', fgColor='1F3864')
THIN = Side(style='thin', color='D9D9D9')


def build(csv_path, out_path):
    nut = nc.load_menu_nutrients(csv_path)
    idx = nc.build_menu_index(nut)
    plan, port = nc.load_plan()
    keys = [k for k in KEY if k in nut.columns]
    forms = list(port['1']['아침'].keys())

    rows = []
    for form in forms:
        for day in range(1, 6):
            for meal in nc.MEALS:
                for r in nc.meal_slots(plan, port, day, meal, form):
                    raw = r['메뉴표기']
                    name = nc.resolve_menu(raw, day, meal)
                    hit = nc.match_name(name, idx)
                    ref, g = r['기준배식량(g)'], r['기본배식량(g)']
                    total = float(nut.at[hit, '총량(g)']) if hit else None
                    factor = nc.serving_factor(r['슬롯'], g, ref, total)
                    rec = {'식사형태': form, '일자': day, '끼니': meal, '슬롯': r['슬롯'],
                           '메뉴': hit or f'[미매칭] {name}',
                           '레시피 총량(g)': total, '1인 제공량(g)': ref, '배식량(g)': g,
                           '배식계수': round(factor, 4) if factor else None,
                           '가정': ''}
                    if hit and factor:
                        for k in keys:
                            rec[k] = float(nut.at[hit, k]) * factor
                    rows.append(rec)
    det = pd.DataFrame(rows)

    wb = openpyxl.Workbook()
    ws_d = wb.create_sheet('슬롯별_상세')
    cols = ['식사형태', '일자', '끼니', '슬롯', '메뉴', '레시피 총량(g)', '1인 제공량(g)', '배식량(g)', '배식계수', '가정'] + keys
    ws_d.append(cols)
    for r in det[cols].itertuples(index=False):
        ws_d.append([None if pd.isna(v) else v for v in r])

    # 요약: 식사형태 × 일자 × 끼니 (SUMIFS)
    ws_s = wb.create_sheet('식사형태별_제공영양량', 0)
    ws_s.append(['식사형태', '일자', '끼니'] + keys)
    combo = det[['식사형태', '일자', '끼니']].drop_duplicates().reset_index(drop=True)
    for i, c in combo.iterrows():
        r = i + 2
        ws_s.cell(r, 1, c['식사형태']); ws_s.cell(r, 2, int(c['일자'])); ws_s.cell(r, 3, c['끼니'])
        for j, k in enumerate(keys):
            col = get_column_letter(cols.index(k) + 1)
            ws_s.cell(r, 4 + j, f"=SUMIFS(슬롯별_상세!{col}:{col},슬롯별_상세!$A:$A,$A{r},"
                                f"슬롯별_상세!$B:$B,$B{r},슬롯별_상세!$C:$C,$C{r})")
    # 요약: 식사형태 × 일자 합계
    ws_y = wb.create_sheet('식사형태별_일자합계', 1)
    ws_y.append(['식사형태', '일자'] + keys)
    combo2 = det[['식사형태', '일자']].drop_duplicates().reset_index(drop=True)
    for i, c in combo2.iterrows():
        r = i + 2
        ws_y.cell(r, 1, c['식사형태']); ws_y.cell(r, 2, int(c['일자']))
        for j, k in enumerate(keys):
            col = get_column_letter(cols.index(k) + 1)
            ws_y.cell(r, 3 + j, f"=SUMIFS(슬롯별_상세!{col}:{col},슬롯별_상세!$A:$A,$A{r},"
                                f"슬롯별_상세!$B:$B,$B{r})")

    ws_n = wb.create_sheet('계산기준')
    for line in [
        ['계산식'],
        ['제공 영양량 = 메뉴 1인분 영양소 × 배식계수'],
        ['배식계수 = 배식량(g) ÷ 1인 제공량(g)'],
        ['섭취 영양소 = 제공 영양량 × 섭취율,  섭취율 = 1 − 잔반등급/4 (0 다먹음 ~ 4 모두 남김)'],
        [],
        ['1인 제공량(분모) 규칙'],
        ['밥/죽 — 밥 계열은 「일반밥/일반찬」의 밥 배식량, 죽·갈죽 계열은 「죽/일반찬」의 죽 배식량'],
        ['김치2(백김치) — 「일반밥/일반찬(백김치)」의 김치2 배식량'],
        ['그 외(국/탕·주찬·부찬1·부찬2·김치1) — 「일반밥/일반찬」의 해당 슬롯 배식량'],
        ['→ 같은 메뉴를 다진찬·갈찬으로 덜 담는 차이가 배식계수에 반영된다'],
        [],
        ['자료 출처'],
        ['영양성분 — CAN Pro 산출 결과 (MemoryResultAll_5일치.csv), 재료별 값을 음식 단위로 합산'],
        ['식단표·배식량 — 조사 앱 NutritionSurveyPage.jsx 의 MEAL_FOODS_BY_DAY / DEFAULT_PORTIONS'],
        [],
        ['참고'],
        ['1·2일차 점심·저녁 「영양죽」은 사용자 확인값으로 배정했습니다 — 1일 점심 계란죽, 1일 저녁 시금치죽, 2일 점심 쇠고기죽, 2일 저녁 버섯죽.'],
        ['CAN Pro 총량(g)은 조리 전 재료 중량 합이므로, 배식량(조리 후 실측)과 단위가 다릅니다.'],
        [f'배식계수 산정 방식: {nc.METHOD}  (A = 기본 배식량을 1인분으로, C = 찬류만 레시피 총량을 1인 제공량으로)'],
    ]:
        ws_n.append(line)

    del wb['Sheet']
    for ws, w2 in ((ws_s, 22), (ws_y, 22), (ws_d, 22), (ws_n, 110)):
        for c in ws[1]:
            c.fill = HDR; c.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=9)
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name=ARIAL, size=10)
                if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith('=')):
                    c.number_format = '#,##0.00'
                c.border = Border(bottom=THIN)
        ws.column_dimensions['A'].width = w2
        for j in range(2, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(j)].width = 13
        if ws is not ws_n:
            ws.freeze_panes = 'D2'
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"
    ws_n.freeze_panes = 'A1'
    for c in ws_n['A']:
        c.font = Font(name=ARIAL, size=10, bold=str(c.value or '').endswith(('계산식', '규칙', '출처', '필요')))
    ws_d.column_dimensions['E'].width = 22
    wb.save(out_path)
    return det


if __name__ == '__main__':
    det = build(sys.argv[1], sys.argv[2])
    print('rows', len(det), '미매칭', det['메뉴'].str.startswith('[미매칭]').sum())
