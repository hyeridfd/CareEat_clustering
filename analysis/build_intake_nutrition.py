# -*- coding: utf-8 -*-
"""
어르신별 섭취 영양소 산출 (5일 × 아침·점심·저녁)

  Supabase(nutrition_survey.meal_portions / plate_waste, elderly_residents.meal_form)
  + CAN Pro 메뉴 영양성분  →  엑셀 3종 시트

실행
  python build_intake_nutrition.py  <CAN Pro CSV 경로>  <출력 xlsx>  [요양원ID]  [A|C]
                                    [--local <elderly_residents.csv> <nutrition_survey.csv>]

  · SUPABASE_URL / SUPABASE_KEY 는 backend/.env 에서 자동으로 읽습니다.
  · 요양원ID를 주면 그 시설만, 생략하면 전체를 계산합니다.
  · 마지막 A/C 는 배식계수 산정 방식 (생략 시 A). 두 방식 설명은 README 참조.
"""
from __future__ import annotations
import json
import sys
import urllib.parse
import urllib.request
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import nutrition_core as nc
from build_serving_nutrition import KEY, ARIAL, HDR, THIN

ENV_CANDIDATES = [HERE.parent / "backend" / ".env", HERE / ".env", HERE.parent / ".env"]


def load_env():
    for p in ENV_CANDIDATES:
        if p.exists():
            env = {}
            for line in p.read_text(encoding="utf-8").splitlines():
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    env[k.strip()] = v.strip().strip('"').strip("'")
            if env.get("SUPABASE_URL") and env.get("SUPABASE_KEY"):
                return env["SUPABASE_URL"].rstrip("/"), env["SUPABASE_KEY"]
    raise SystemExit("SUPABASE_URL / SUPABASE_KEY 를 찾지 못했습니다 (backend/.env 확인)")


def fetch(url, key, table, select="*", eq=None, page=1000):
    rows, offset = [], 0
    while True:
        q = {"select": select, "limit": str(page), "offset": str(offset)}
        if eq:
            q.update({k: f"eq.{v}" for k, v in eq.items()})
        req = urllib.request.Request(f"{url}/rest/v1/{table}?" + urllib.parse.urlencode(q),
                                     headers={"apikey": key, "Authorization": f"Bearer {key}"})
        got = json.load(urllib.request.urlopen(req, timeout=60))
        rows += got
        if len(got) < page:
            return rows
        offset += page


def as_dict(v):
    """문자열로 두 번 감싸인 JSON도 풀어낸다."""
    for _ in range(3):
        if isinstance(v, dict):
            return v
        if not isinstance(v, str):
            return {}
        try:
            v = json.loads(v)
        except Exception:
            return {}
    return v if isinstance(v, dict) else {}


def from_csv(residents_csv, survey_csv):
    """Supabase 대신 내려받은 CSV 두 개로 계산할 때 쓴다."""
    def read(path):
        for enc in ("utf-8-sig", "cp949"):
            try:
                return pd.read_csv(path, encoding=enc)
            except UnicodeDecodeError:
                continue
        raise SystemExit(f"인코딩을 알 수 없습니다: {path}")
    res = read(residents_csv).to_dict("records")
    sv = read(survey_csv).to_dict("records")
    return res, sv


def build(csv_path, out_path, home=None, method=None, csv_source=None):
    if method:
        nc.METHOD = method
    nut = nc.load_menu_nutrients(csv_path)
    idx = nc.build_menu_index(nut)
    plan, port = nc.load_plan()
    keys = [k for k in KEY if k in nut.columns]

    if csv_source:
        res_rows, surveys = from_csv(*csv_source)
    else:
        url, key = load_env()
        eq = {"nursing_home_id": home} if home else None
        res_rows = fetch(url, key, "elderly_residents", "id,name,nursing_home_id,meal_form", eq)
        surveys = fetch(url, key, "nutrition_survey",
                        "elderly_id,nursing_home_id,meal_portions,plate_waste,updated_at", eq)
    if home:
        res_rows = [r for r in res_rows if r.get("nursing_home_id") == home]
        surveys = [s for s in surveys if s.get("nursing_home_id") == home]
    residents = {r["id"]: r for r in res_rows}
    # 계산 대상은 영양조사가 있는 어르신뿐
    print(f"어르신 {len(residents)}명 · 영양조사 {len(surveys)}건")

    rows = []
    for s in surveys:
        eid = s["elderly_id"]
        res = residents.get(eid, {})
        form = nc.NAME_ALIAS.get(res.get("meal_form"), res.get("meal_form")) or "일반밥/일반찬"
        if form not in port["1"]["아침"]:
            form = "일반밥/일반찬"
        mp, pw = as_dict(s.get("meal_portions")), as_dict(s.get("plate_waste"))
        for day in range(1, 6):
            for meal in nc.MEALS:
                served = (mp.get(f"day{day}") or {}).get(meal) or {}
                waste = (pw.get(f"day{day}") or {}).get(meal) or {}
                if not served:
                    continue
                for r in nc.meal_slots(plan, port, day, meal, form, elderly_id=eid):
                    slot = r["슬롯"]
                    g = served.get(slot, r["기본배식량(g)"])
                    name = nc.resolve_menu(r["메뉴표기"], day, meal)
                    hit = nc.match_name(name, idx)
                    deferred = isinstance(g, str)
                    g = None if deferred else g
                    total = float(nut.at[hit, "총량(g)"]) if hit else None
                    factor = nc.serving_factor(slot, g, r["기준배식량(g)"], total)
                    entered = slot in waste
                    rate = nc.intake_rate(waste.get(slot, 0))
                    rec = {"어르신ID": eid, "성명": res.get("name"), "요양원": res.get("nursing_home_id"),
                           "식사형태": form, "일자": day, "끼니": meal, "슬롯": slot,
                           "메뉴": hit or f"[미매칭] {name}",
                           "배식량(g)": g, "1인 제공량(g)": r["기준배식량(g)"],
                           "레시피 총량(g)": total,
                           "배식계수": round(factor, 4) if factor else None,
                           "잔반등급": waste.get(slot, 0), "섭취율": rate,
                           "잔반입력": "입력" if entered else "기본(다 먹음)",
                           "비고": "추후섭취" if deferred else ""}
                    if hit and factor and rate is not None:
                        for k in keys:
                            rec[k] = float(nut.at[hit, k]) * factor * rate
                    rows.append(rec)

    det = pd.DataFrame(rows)
    if det.empty:
        raise SystemExit("계산할 조사 데이터가 없습니다.")

    wb = openpyxl.Workbook()
    head = ["어르신ID", "성명", "요양원", "식사형태", "일자", "끼니", "슬롯", "메뉴",
            "레시피 총량(g)", "1인 제공량(g)", "배식량(g)", "배식계수", "잔반등급", "잔반입력", "섭취율", "비고"]
    ws_d = wb.create_sheet("슬롯별_상세")
    ws_d.append(head + keys)
    for r in det[head + keys].itertuples(index=False):
        ws_d.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])

    def summary(sheet, by, pos):
        ws = wb.create_sheet(sheet, pos)
        ws.append(by + keys)
        combo = det[by].drop_duplicates().reset_index(drop=True)
        for i, c in combo.iterrows():
            r = i + 2
            for j, b in enumerate(by):
                ws.cell(r, j + 1, c[b] if not pd.isna(c[b]) else None)
            conds = "".join(f",슬롯별_상세!${get_column_letter(head.index(b) + 1)}:"
                            f"${get_column_letter(head.index(b) + 1)},${get_column_letter(j + 1)}{r}"
                            for j, b in enumerate(by))
            for j, k in enumerate(keys):
                col = get_column_letter(len(head) + 1 + j)
                ws.cell(r, len(by) + 1 + j,
                        f"=SUMIFS(슬롯별_상세!{col}:{col}{conds})")
        return ws

    ws_m = summary("어르신_일자_끼니", ["어르신ID", "성명", "식사형태", "일자", "끼니"], 0)
    ws_y = summary("어르신_일자합계", ["어르신ID", "성명", "식사형태", "일자"], 1)

    ws_n = wb.create_sheet("계산기준")
    for line in [
        ["섭취 영양소 = 메뉴 1인분 영양소 × 배식계수 × 섭취율"],
        ["배식계수 = 실제 배식량(g) ÷ 1인 제공량(g)"],
        ["섭취율 = 1 − 잔반등급/4   (0 다먹음 · 1 25% · 2 50% · 3 75% · 4 모두 남김)"],
        [f"배식계수 산정 방식: {nc.METHOD}  (A = 기본 배식량을 1인분으로, C = 찬류만 레시피 총량을 1인 제공량으로)"],
        [],
        ["배식량은 조사에 입력된 값을 쓰고, 입력이 없으면 식사형태별 기본 배식량을 씁니다."],
        ["'추후섭취'로 표시된 칸은 해당 끼니 섭취에서 제외했습니다 (비고 열)."],
        ["잔반이 입력되지 않은 칸은 조사 화면의 기본값대로 '다 먹음(0)'으로 계산했습니다 ('잔반입력' 열에 표시)."],
        ["4·5일차 밥/죽은 HS09·HS29 어르신만 흰죽으로 대체 계산합니다 (식단표 표기 반영)."],
        ["1·2일차 점심·저녁 영양죽 = 계란죽 / 시금치죽 / 쇠고기죽 / 버섯죽"],
    ]:
        ws_n.append(line)

    del wb["Sheet"]
    for ws in (ws_m, ws_y, ws_d, ws_n):
        for c in ws[1]:
            c.fill = HDR; c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=9)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name=ARIAL, size=10)
                if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                    c.number_format = "#,##0.00"
                c.border = Border(bottom=THIN)
        for j in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(j)].width = 14
        if ws is not ws_n:
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"
    ws_n.column_dimensions["A"].width = 110
    wb.save(out_path)
    print("저장:", out_path, "| 상세", len(det), "행")
    return det


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    args = sys.argv[3:]
    csv_source = None
    if "--local" in args:
        i = args.index("--local")
        csv_source = (args[i + 1], args[i + 2])
        args = args[:i] + args[i + 3:]
    method = next((a.upper() for a in args if a.upper() in ("A", "C")), None)
    home = next((a for a in args if a.upper() not in ("A", "C")), None)
    build(sys.argv[1], sys.argv[2], home, method, csv_source)
