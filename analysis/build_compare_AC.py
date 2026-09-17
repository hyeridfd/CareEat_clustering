# -*- coding: utf-8 -*-
"""배식계수 산정 방식 A·C 비교 — 식사형태별 제공 영양량(잔반 반영 전)"""
import sys
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import nutrition_core as nc
from build_serving_nutrition import KEY, ARIAL, HDR, THIN

CMP = ['에너지(kcal)', '단백질(g)', '지질(g)', '탄수화물(g)', '총 식이섬유(g)',
       '나트륨(mg)', '칼슘(mg)', '철(mg)', '칼륨(mg)', '비타민 C(mg)']


def detail(nut, idx, plan, port, keys, method):
    rows = []
    for form in port['1']['아침']:
        for day in range(1, 6):
            for meal in nc.MEALS:
                for r in nc.meal_slots(plan, port, day, meal, form):
                    hit = nc.match_name(nc.resolve_menu(r['메뉴표기'], day, meal), idx)
                    total = float(nut.at[hit, '총량(g)']) if hit else None
                    f = nc.serving_factor(r['슬롯'], r['기본배식량(g)'], r['기준배식량(g)'], total, method)
                    rec = {'식사형태': form, '일자': day, '끼니': meal, '슬롯': r['슬롯'], '메뉴': hit,
                           '레시피 총량(g)': total, '1인 제공량(g)': r['기준배식량(g)'],
                           '배식량(g)': r['기본배식량(g)'], '배식계수': round(f, 4) if f else None}
                    if hit and f:
                        for k in keys:
                            rec[k] = float(nut.at[hit, k]) * f
                    rows.append(rec)
    return pd.DataFrame(rows)


def style(ws, note=False):
    for c in ws[1]:
        c.fill = HDR; c.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=9)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=ARIAL, size=10)
            if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith('=')):
                c.number_format = '#,##0.00'
            c.border = Border(bottom=THIN)
    for j in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.column_dimensions['A'].width = 110 if note else 22
    if not note:
        ws.freeze_panes = 'D2'
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"


def main(csv_path, out_path):
    nut = nc.load_menu_nutrients(csv_path)
    idx = nc.build_menu_index(nut)
    plan, port = nc.load_plan()
    keys = [k for k in KEY if k in nut.columns]
    dets = {m: detail(nut, idx, plan, port, keys, m) for m in ('A', 'C')}

    wb = openpyxl.Workbook()
    head = ['식사형태', '일자', '끼니', '슬롯', '메뉴', '레시피 총량(g)', '1인 제공량(g)', '배식량(g)', '배식계수']
    sheets = {}
    for m in ('A', 'C'):
        ws = wb.create_sheet(f'{m}_슬롯별_상세')
        ws.append(head + keys)
        for r in dets[m][head + keys].itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
        sheets[m] = ws
        style(ws)

    for m in ('A', 'C'):                      # 끼니별 요약
        ws = wb.create_sheet(f'{m}_끼니별', 0 if m == 'A' else 1)
        ws.append(['식사형태', '일자', '끼니'] + keys)
        combo = dets[m][['식사형태', '일자', '끼니']].drop_duplicates().reset_index(drop=True)
        for i, c in combo.iterrows():
            r = i + 2
            ws.cell(r, 1, c['식사형태']); ws.cell(r, 2, int(c['일자'])); ws.cell(r, 3, c['끼니'])
            for j, k in enumerate(keys):
                col = get_column_letter(len(head) + 1 + j)
                ws.cell(r, 4 + j, f"=SUMIFS({m}_슬롯별_상세!{col}:{col},{m}_슬롯별_상세!$A:$A,$A{r},"
                                  f"{m}_슬롯별_상세!$B:$B,$B{r},{m}_슬롯별_상세!$C:$C,$C{r})")
        style(ws)

    ws = wb.create_sheet('비교_일자합계', 0)   # A·C 나란히
    cols = ['식사형태', '일자']
    for k in CMP:
        cols += [f'A · {k}', f'C · {k}', f'C/A · {k}']
    ws.append(cols)
    combo = dets['A'][['식사형태', '일자']].drop_duplicates().reset_index(drop=True)
    for i, c in combo.iterrows():
        r = i + 2
        ws.cell(r, 1, c['식사형태']); ws.cell(r, 2, int(c['일자']))
        for j, k in enumerate(CMP):
            col = get_column_letter(len(head) + 1 + keys.index(k))
            a = get_column_letter(3 + j * 3); cc = get_column_letter(4 + j * 3)
            ws.cell(r, 3 + j * 3, f"=SUMIFS(A_슬롯별_상세!{col}:{col},A_슬롯별_상세!$A:$A,$A{r},A_슬롯별_상세!$B:$B,$B{r})")
            ws.cell(r, 4 + j * 3, f"=SUMIFS(C_슬롯별_상세!{col}:{col},C_슬롯별_상세!$A:$A,$A{r},C_슬롯별_상세!$B:$B,$B{r})")
            ws.cell(r, 5 + j * 3, f"=IFERROR({cc}{r}/{a}{r},\"\")").number_format = '0.0%'
    style(ws)
    ws.freeze_panes = 'C2'

    wsn = wb.create_sheet('계산기준')
    for line in [
        ['제공 영양량 = 메뉴 1인분 영양소 × 배식계수   (섭취 영양소는 여기에 섭취율 = 1 − 잔반/4 를 곱한다)'],
        [],
        ['A — 각 식사형태의 기본 배식량을 그 메뉴 1인분으로 본다. 배식계수 = 배식량 ÷ 기준 배식량'],
        ['     기준 배식량: 밥은 「일반밥/일반찬」의 밥, 죽·갈죽은 「죽/일반찬」의 죽,'],
        ['     김치2는 「일반밥/일반찬(백김치)」의 김치2, 나머지는 「일반밥/일반찬」의 해당 슬롯'],
        [],
        ['C — 밥/죽·국/탕은 A와 동일. 주찬·부찬·김치는 배식계수 = 배식량 ÷ 레시피 총량'],
        ['     찬류는 조리 수율이 1에 가까워 배식량 ÷ 레시피 총량이 실제 제공 비율에 가깝다.'],
        ['     밥/죽은 생쌀 기준, 국은 물이 재료에 없어 레시피 총량을 분모로 쓸 수 없다.'],
        [],
        ['자료 — 영양성분: CAN Pro 산출 결과(MemoryResultAll_5일치.csv), 재료별 값을 음식 단위로 합산'],
        ['       식단표·배식량: 조사 앱 NutritionSurveyPage.jsx 의 MEAL_FOODS_BY_DAY / DEFAULT_PORTIONS'],
        ['       1·2일차 점심·저녁 영양죽 = 계란죽 / 시금치죽 / 쇠고기죽 / 버섯죽 (사용자 확인)'],
        ['       간식은 CAN Pro 자료가 없어 제외 (아침·점심·저녁만)'],
    ]:
        wsn.append(line)
    style(wsn, note=True)
    del wb['Sheet']
    wb.save(out_path)
    print('saved', out_path)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
